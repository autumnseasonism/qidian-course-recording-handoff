"""Export explicitly selected candidate frames. Does not certify completeness or human review."""
import argparse
import csv
import hashlib
import io
import json
import math
import subprocess
from pathlib import Path

from PIL import Image
from reportlab.pdfgen import canvas
from reportlab.lib.utils import ImageReader
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from pypdf import PdfReader
from pptx import Presentation
from pptx.util import Inches, Pt

from scan_pages import guarded_path, sha, file_identity


def finite(value):
    return type(value) in (float, int) and math.isfinite(value)


def capture(path):
    before = file_identity(path)
    digest = sha(path)
    if before != file_identity(path):
        raise ValueError('Input changed while hashing: ' + str(path))
    return (path, before, digest)


def verify(records, root):
    for path, identity, digest in records:
        guarded_path(path, root, must_exist=True)
        if file_identity(path) != identity or sha(path) != digest or file_identity(path) != identity:
            raise ValueError('Input changed: ' + str(path))


def main():
    ap = argparse.ArgumentParser(description=__doc__)
    for name in ('workspace-root', 'candidates', 'selection', 'output-dir', 'ffprobe'):
        ap.add_argument('--' + name, required=True)
    ap.add_argument('--font-file', help='Explicit installed TTF/TTC with required title glyphs; required for non-ASCII titles')
    args = ap.parse_args()
    root = guarded_path(args.workspace_root, must_exist=True)
    if not root.is_dir():
        raise ValueError('Workspace root must be a directory')
    index = guarded_path(args.candidates, root, must_exist=True)
    selection_path = guarded_path(args.selection, root, must_exist=True)
    out = guarded_path(args.output_dir, root)
    if out.exists() or not index.is_file() or not selection_path.is_file():
        raise ValueError('Inputs must be regular files; output directory must be new')
    if any((index.parent / marker).exists() for marker in ('FAILED.json', '候选索引发布失败.json')):
        raise ValueError('Candidate directory has a failure marker; it cannot be exported')
    records = [capture(index), capture(selection_path)]
    data = json.loads(index.read_text(encoding='utf-8-sig'))
    selection = json.loads(selection_path.read_text(encoding='utf-8-sig'))
    signature = data['signature']
    source = guarded_path(signature['input'], root, must_exist=True)
    source_record = capture(source)
    records.append(source_record)
    if source_record[1][2] != signature['input_bytes'] or source_record[2] != signature['input_sha256']:
        raise ValueError('Source differs from candidate signature')
    duration = signature['duration_seconds']
    if not finite(duration) or duration <= 0:
        raise ValueError('Invalid source duration')
    ffprobe = guarded_path(args.ffprobe, must_exist=True)
    if not ffprobe.is_file():
        raise ValueError('FFprobe must be a regular executable file')
    info = json.loads(subprocess.run([str(ffprobe), '-v', 'error', '-show_streams', '-show_format',
                                     '-of', 'json', str(source)], capture_output=True, check=True).stdout)
    video = next(item for item in info['streams'] if item.get('codec_type') == 'video')
    actual_duration = float(info['format']['duration'])
    if not math.isfinite(actual_duration) or abs(actual_duration-duration) > 0.000001:
        raise ValueError('Actual duration differs from candidate signature')
    if (video['width'], video['height']) != (signature['width'], signature['height']):
        raise ValueError('Actual video dimensions differ from candidate signature')
    if selection.get('candidate_index_sha256') != records[0][2] or selection.get('source_sha256') != source_record[2]:
        raise ValueError('Selection must bind exact candidate index and source hashes')
    candidates = data['candidates']
    if not isinstance(candidates, list) or not candidates:
        raise ValueError('Empty candidates')
    lookup = {}
    for item in candidates:
        identifier = item['index']
        if type(identifier) is not int or identifier < 1 or identifier in lookup:
            raise ValueError('Duplicate/invalid candidate ID')
        timestamp = item['actual_seconds']
        if not finite(timestamp) or not 0 <= timestamp < duration:
            raise ValueError('Invalid candidate timestamp')
        lookup[identifier] = item
    requested = selection.get('pages')
    if not isinstance(requested, list) or not requested:
        raise ValueError('Explicit nonempty page selection is required')
    pages, image_bytes, previous_last, previous_rep = [], [], -1, -1
    for number, request in enumerate(requested, 1):
        ids = [request.get(key) for key in ('first_candidate', 'representative_candidate', 'last_candidate')]
        if any(type(value) is not int or value not in lookup for value in ids):
            raise ValueError('Selected candidate IDs must exist')
        first, representative, last = [lookup[value] for value in ids]
        times = [item['actual_seconds'] for item in (first, representative, last)]
        if not (times[0] <= times[1] <= times[2] and times[0] >= previous_last and times[1] > previous_rep):
            raise ValueError('Pages must follow source order without overlap/repeated representative frames')
        previous_last, previous_rep = times[2], times[1]
        title = request.get('title')
        if not isinstance(title, str) or not title.strip():
            raise ValueError('Every page needs an explicit title')
        raw_relative = Path(representative['image'])
        if raw_relative.is_absolute() or '..' in raw_relative.parts:
            raise ValueError('Candidate image must be a safe relative path')
        image = guarded_path(index.parent / raw_relative, root, must_exist=True)
        if index.parent not in image.parents or not image.is_file():
            raise ValueError('Image must be a regular file inside candidate directory')
        record = capture(image)
        records.append(record)
        if record[2] != representative['sha256']:
            raise ValueError('Candidate image hash mismatch')
        payload = image.read_bytes()
        with Image.open(io.BytesIO(payload)) as check:
            if check.format != 'PNG':
                raise ValueError('Candidate must be PNG')
            check.load()
            size = check.size
        if size != (signature['width'], signature['height']):
            raise ValueError('Candidate dimensions differ from complete video frame')
        image_bytes.append(payload)
        pages.append(dict(number=number, title=title, first_seconds=times[0],
                          representative_seconds=times[1], last_seconds=times[2],
                          first_candidate=ids[0], representative_candidate=ids[1], last_candidate=ids[2],
                          image=f'pages/{number:04d}.png', image_sha256=record[2],
                          width=size[0], height=size[1],
                          animation_note=request.get('animation_note', 'unreviewed')))
    title_font = 'Helvetica'
    font_record = None
    if args.font_file:
        font_path = guarded_path(args.font_file, must_exist=True)
        if not font_path.is_file() or font_path.suffix.lower() not in ('.ttf', '.ttc'):
            raise ValueError('Explicit title font must be an installed TTF/TTC file')
        font_record = capture(font_path)
        font = TTFont('CoursePageTitle', str(font_path))
        pdfmetrics.registerFont(font)
        for page in pages:
            if any(ord(char) not in font.face.charToGlyph for char in page['title'] if not char.isspace()):
                raise ValueError('Title font is missing required glyphs')
        title_font = 'CoursePageTitle'
    elif any(not page['title'].isascii() for page in pages):
        raise ValueError('Non-ASCII titles require explicit --font-file; no font is installed or copied automatically')
    title_lines = []
    for page in pages:
        if any(ord(char) < 32 and char != '\n' for char in page['title']):
            raise ValueError('Title contains unsupported control characters')
        lines = []
        for paragraph in page['title'].split('\n'):
            line = ''
            for char in paragraph:
                if pdfmetrics.stringWidth(line + char, title_font, 14) > 1240:
                    lines.append(line)
                    line = char
                else:
                    line += char
            lines.append(line)
        if len(lines) > 5:
            raise ValueError('Complete title cannot fit five readable footer lines; shorten the explicitly selected title, not silently truncate')
        title_lines.append(lines)

    def verify_all():
        verify(records, root)
        if font_record:
            font_path, identity, digest = font_record
            guarded_path(font_path, must_exist=True)
            if file_identity(font_path) != identity or sha(font_path) != digest or file_identity(font_path) != identity:
                raise ValueError('Installed font changed during export')
        if any((index.parent / marker).exists() for marker in ('FAILED.json', '候选索引发布失败.json')):
            raise ValueError('Candidate failure marker appeared during export')

    verify_all()
    guarded_path(out, root)
    out.mkdir(parents=True, exist_ok=False)
    (out / 'pages').mkdir()
    try:
        for page, payload in zip(pages, image_bytes):
            with guarded_path(out / page['image'], root).open('xb') as stream:
                stream.write(payload)
        pdf = guarded_path(out / 'pages.pdf', root)
        if pdf.exists():
            raise ValueError('PDF output appeared concurrently')
        with pdf.open('xb') as pdf_stream:
            pdf_canvas = canvas.Canvas(pdf_stream, pagesize=(1280, 840))
            for page, payload, lines in zip(pages, image_bytes, title_lines):
                pdf_canvas.drawImage(ImageReader(io.BytesIO(payload)), 0, 120, width=1280, height=720,
                                     preserveAspectRatio=True, anchor='c', mask='auto')
                pdf_canvas.setFont(title_font, 14)
                for line_number, line in enumerate(lines):
                    pdf_canvas.drawString(20, 100 - line_number * 18, line)
                pdf_canvas.setFont('Helvetica', 12)
                pdf_canvas.drawString(20, 15, f"{page['number']}/{len(pages)} | video {page['first_seconds']:.3f}-{page['last_seconds']:.3f}s | frame {page['representative_seconds']:.3f}s")
                pdf_canvas.showPage()
            pdf_canvas.save()
        reader = PdfReader(pdf)
        if len(reader.pages) != len(pages):
            raise ValueError('Actual PDF page count mismatch')
        # Compare decoded pixels because PDF lossless image encoding changes bytes.
        pdf_pixels = []
        for pdf_page, payload in zip(reader.pages, image_bytes):
            embedded = list(pdf_page.images)
            if len(embedded) != 1:
                raise ValueError('PDF must embed exactly one whole-frame image per page')
            with Image.open(io.BytesIO(payload)) as original:
                expected_pixels = hashlib.sha256(original.convert('RGB').tobytes()).hexdigest()
            embedded_image = embedded[0].image.convert('RGB')
            actual_pixels = hashlib.sha256(embedded_image.tobytes()).hexdigest()
            if actual_pixels != expected_pixels or embedded_image.size != (pages[len(pdf_pixels)]['width'], pages[len(pdf_pixels)]['height']):
                raise ValueError('PDF embedded pixels differ from source')
            pdf_pixels.append(actual_pixels)
        ppt = Presentation()
        ppt.slide_width, ppt.slide_height = Inches(12.8), Inches(7.6)
        for page, payload in zip(pages, image_bytes):
            slide = ppt.slides.add_slide(ppt.slide_layouts[6])
            scale = min(12.8 / page['width'], 7.2 / page['height'])
            width, height = page['width'] * scale, page['height'] * scale
            slide.shapes.add_picture(io.BytesIO(payload), Inches((12.8-width)/2), Inches((7.2-height)/2), Inches(width), Inches(height))
            text = slide.shapes.add_textbox(Inches(.2), Inches(7.2), Inches(12.4), Inches(.35)).text_frame
            text.text = f"{page['number']}/{len(pages)} | video {page['first_seconds']:.3f}-{page['last_seconds']:.3f}s | frame {page['representative_seconds']:.3f}s"
            text.paragraphs[0].font.size = Pt(12)
            slide.notes_slide.notes_text_frame.text = page['title']
        ppt_path = guarded_path(out / 'pages.pptx', root)
        if ppt_path.exists():
            raise ValueError('PPT output appeared concurrently')
        with ppt_path.open('xb') as ppt_stream:
            ppt.save(ppt_stream)
        opened = Presentation(ppt_path)
        if len(opened.slides) != len(pages):
            raise ValueError('Actual PPT slide count mismatch')
        for slide, page in zip(opened.slides, pages):
            pictures = [shape for shape in slide.shapes if shape.shape_type == 13]
            if len(pictures) != 1 or hashlib.sha256(pictures[0].image.blob).hexdigest() != page['image_sha256']:
                raise ValueError('PPT embedded image SHA256 mismatch')
            if slide.notes_slide.notes_text_frame.text != page['title']:
                raise ValueError('PPT full page title notes mismatch')
        verify_all()
        for page in pages:
            if sha(guarded_path(out / page['image'], root, must_exist=True)) != page['image_sha256']:
                raise ValueError('Exported PNG hash mismatch')
        result = dict(schema_version=1, signature=signature, pages=pages,
                      candidate_index_sha256=records[0][2], selection_sha256=records[1][2],
                      status='machine_checked_human_review_pending', completeness_verified=False,
                      human_review_verified=False, human_review_claim=selection.get('review', {}),
                      pdf=dict(file=pdf.name, sha256=sha(pdf), pages=len(reader.pages), embedded_rgb_sha256=pdf_pixels),
                      pptx=dict(file=ppt_path.name, sha256=sha(ppt_path), slides=len(opened.slides), embedded_png_sha256_verified=True),
                      title_font=dict(file=str(font_record[0]), sha256=font_record[2]) if font_record else dict(name='Helvetica'),
                      rendering_status='not_rendered', limitations=['Candidate scanning may miss small changes or brief animation states', 'Selection is explicit but not proof of completeness', 'PDF and PPT require rendering and visual inspection'])
        with guarded_path(out / 'pages.csv', root).open('x', encoding='utf-8-sig', newline='') as stream:
            writer = csv.DictWriter(stream, fieldnames=list(pages[0]))
            writer.writeheader()
            writer.writerows(pages)
        result['outputs'] = {page['image']: page['image_sha256'] for page in pages}
        for artifact in (pdf, ppt_path, out / 'pages.csv'):
            result['outputs'][artifact.name] = sha(artifact)
        verify_all()
        with guarded_path(out / 'page-index.json', root).open('x', encoding='utf-8') as stream:
            json.dump(result, stream, ensure_ascii=False, indent=2)
        verify_all()
    except BaseException as error:
        # Preserve incomplete diagnostics; never label a failed run as accepted.
        # Failure marker takes precedence over any already written index. Keep
        # files intact rather than delete a concurrently changed file.
        guarded_path(out, root, must_exist=True)
        with (out / 'FAILED.json').open('x', encoding='utf-8') as stream:
            json.dump(dict(status='failed', error=str(error)), stream, ensure_ascii=False, indent=2)
        raise
    print(json.dumps(dict(output=str(out), pages=len(pages), status=result['status'])))


if __name__ == '__main__':
    main()
